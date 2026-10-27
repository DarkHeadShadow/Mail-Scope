"""Attachment preview renderers; these functions run inside the isolated worker."""

__author__ = "Bálint Szabolcs"

import csv
import io
from typing import Any

from sources.content import html_to_text


TEXT_LIMIT_BYTES = 5 * 1024 * 1024
TABLE_PAGE_ROWS = 2000
IMAGE_MAX_SIDE = 4096


def _decode(data: bytes) -> tuple[str, str]:
    """Decode text with BOM detection, strict UTF-8, then Windows-1250 and Latin-1."""
    for bom, codec in ((b"\xef\xbb\xbf", "utf-8-sig"), (b"\xff\xfe", "utf-16"), (b"\xfe\xff", "utf-16")):
        if data.startswith(bom):
            return data.decode(codec, errors="replace"), codec.upper()
    for codec, label in (("utf-8", "UTF-8"), ("cp1250", "Windows-1250")):
        try:
            return data.decode(codec), label
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1"), "ISO-8859-1"


def render_text(data: bytes, filename: str) -> dict[str, Any]:
    """Return text (and for HTML, the cleaned text plus the source) up to 5 MB."""
    truncated = len(data) > TEXT_LIMIT_BYTES
    text, encoding = _decode(data[:TEXT_LIMIT_BYTES])
    lowered = filename.casefold()
    is_html = lowered.endswith((".html", ".htm")) or text.lstrip()[:200].casefold().startswith(
        ("<!doctype html", "<html")
    )
    result: dict[str, Any] = {
        "kind": "text",
        "encoding": encoding,
        "truncated": truncated,
        "text": html_to_text(text) if is_html else text,
    }
    if is_html:
        result["source"] = text
    return result


def render_table(data: bytes, page: int = 0) -> dict[str, Any]:
    """Return one page of up to 2000 CSV rows with the detected delimiter."""
    text, encoding = _decode(data[:TEXT_LIMIT_BYTES * 4])
    sample = text[:20000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";" if sample.count(";") > sample.count(",") else ","
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    start = max(0, page) * TABLE_PAGE_ROWS
    return {
        "kind": "table",
        "encoding": encoding,
        "delimiter": {"\t": "tabulátor"}.get(delimiter, delimiter),
        "total_rows": len(rows),
        "page": max(0, page),
        "page_size": TABLE_PAGE_ROWS,
        "rows": rows[start:start + TABLE_PAGE_ROWS],
    }


def render_image(data: bytes, page: int = 0) -> dict[str, Any]:
    """Re-encode one image page as PNG and read EXIF and GPS data.

    The browser only receives this re-encoded PNG, never the original bytes.
    Pillow's decompression-bomb guard stays active.
    """
    from PIL import ExifTags, Image, ImageOps

    with Image.open(io.BytesIO(data)) as image:
        source_format = image.format or ""
        pages = getattr(image, "n_frames", 1)
        image.seek(min(max(0, page), pages - 1))
        width, height = image.size
        exif = image.getexif()
        details: dict[str, str] = {}
        for tag, value in exif.items():
            name = ExifTags.TAGS.get(tag)
            if name in {"Make", "Model", "DateTime", "Software", "Artist", "Copyright", "ImageDescription"}:
                details[name] = str(value).strip("\x00 ")
        try:
            sub = exif.get_ifd(ExifTags.IFD.Exif)
            for tag in (0x9003, 0x9004):  # DateTimeOriginal, DateTimeDigitized
                if tag in sub:
                    details[ExifTags.TAGS.get(tag, str(tag))] = str(sub[tag]).strip("\x00 ")
        except Exception:
            pass
        gps = _gps_text(exif)
        frame = ImageOps.exif_transpose(image.copy())
        if frame.mode not in {"RGB", "RGBA", "L", "LA"}:
            frame = frame.convert("RGBA" if "A" in frame.getbands() else "RGB")
        frame.thumbnail((IMAGE_MAX_SIDE, IMAGE_MAX_SIDE))
        output = io.BytesIO()
        frame.save(output, format="PNG")
    return {
        "kind": "image",
        "png": output.getvalue(),
        "width": width,
        "height": height,
        "pages": pages,
        "page": min(max(0, page), pages - 1),
        "format": Image.MIME.get(source_format, source_format),
        "exif": details,
        "gps": gps,
    }


def _gps_text(exif: Any) -> str | None:
    """Render EXIF GPS coordinates as plain decimal text (no map is shown)."""
    try:
        from PIL import ExifTags

        gps = exif.get_ifd(ExifTags.IFD.GPSInfo)
    except Exception:
        return None
    if not gps or 2 not in gps or 4 not in gps:
        return None

    def decimal(values: Any, reference: str) -> float:
        degrees, minutes, seconds = (float(value) for value in values)
        result = degrees + minutes / 60 + seconds / 3600
        return -result if reference in {"S", "W"} else result

    try:
        latitude = decimal(gps[2], str(gps.get(1, "N")))
        longitude = decimal(gps[4], str(gps.get(3, "E")))
    except (TypeError, ValueError, ZeroDivisionError):
        return None
    return f"{latitude:.6f}, {longitude:.6f}"


# --- Stage 2: documents -----------------------------------------------------

DOCUMENT_TEXT_LIMIT = 2 * 1024 * 1024
SHEET_PAGE_ROWS = 2000
SHEET_MAX_COLUMNS = 200
_ALLOWED_TAGS = {
    "p", "br", "hr", "div", "span", "h1", "h2", "h3", "h4", "h5", "h6", "strong", "b", "em", "i",
    "u", "s", "sub", "sup", "small", "blockquote", "pre", "code", "ul", "ol", "li", "dl", "dt", "dd",
    "table", "thead", "tbody", "tfoot", "tr", "td", "th", "caption", "colgroup", "col", "img", "a",
    "figure", "figcaption", "center", "font",
}
_DOCUMENT_STYLE = (
    "body{margin:24px;color:#1b1f1d;background:#fff;font:14px/1.6 'Segoe UI',Arial,sans-serif}"
    "table{border-collapse:collapse;margin:10px 0}td,th{border:1px solid #bbb;padding:4px 8px}"
    "img{max-width:100%;height:auto}a{color:#1a6b42;text-decoration:underline dotted}"
)


def sanitize_html(markup: str, *, keep_styles: bool) -> str:
    """Remove scripts, event handlers, link targets and every non-data URL.

    The result is shown in a sandboxed iframe with its own strict CSP, so
    even a missed construct can neither run nor load anything external.
    """
    import nh3

    attributes = {
        "*": {"title", "lang", "dir", "align", "colspan", "rowspan", "width", "height"} | ({"style"} if keep_styles else set()),
        "img": {"src", "alt"},
        "font": {"color", "size", "face"},
        "td": {"colspan", "rowspan", "align", "valign"},
        "th": {"colspan", "rowspan", "align", "valign"},
    }
    return nh3.clean(
        markup,
        tags=_ALLOWED_TAGS | ({"style"} if keep_styles else set()),
        clean_content_tags={"script", "noscript", "iframe", "object", "embed", "template", "title"},
        attributes=attributes,
        url_schemes={"data"},
        link_rel=None,
        strip_comments=True,
    )


def _document(body: str, title: str) -> str:
    """Wrap sanitized markup into a minimal standalone HTML document."""
    import html

    return (
        "<!doctype html><html lang=\"hu\"><head><meta charset=\"utf-8\">"
        f"<title>{html.escape(title)}</title><style>{_DOCUMENT_STYLE}</style></head>"
        f"<body>{body}</body></html>"
    )


def _limited(text: str) -> tuple[str, bool]:
    """Cut extracted text at the 2 MB document text limit."""
    encoded = text.encode("utf-8")
    if len(encoded) <= DOCUMENT_TEXT_LIMIT:
        return text, False
    return encoded[:DOCUMENT_TEXT_LIMIT].decode("utf-8", errors="ignore"), True


def render_html_document(data: bytes, filename: str) -> dict[str, Any]:
    """Return a sanitized HTML document for the sandboxed formatted view."""
    text, _ = _decode(data[:TEXT_LIMIT_BYTES])
    return {"kind": "html", "document": _document(sanitize_html(text, keep_styles=True), filename)}


def render_pdf(data: bytes, page: int = 0, scale: float = 2.0) -> dict[str, Any]:
    """Render one PDF page to PNG with its text layer; no JavaScript or forms run."""
    import pypdfium2 as pdfium

    document = pdfium.PdfDocument(data)
    try:
        pages = len(document)
        index = min(max(0, page), max(0, pages - 1))
        pdf_page = document[index]
        width, height = pdf_page.get_size()
        bitmap = pdf_page.render(scale=scale)
        image = bitmap.to_pil()
        output = io.BytesIO()
        image.save(output, format="PNG")
        text = pdf_page.get_textpage().get_text_range()
    finally:
        document.close()
    return {
        "kind": "pdf",
        "png": output.getvalue(),
        "page": index,
        "pages": pages,
        "width": round(width * scale),
        "height": round(height * scale),
        "text": text,
    }


def _pdf_signals(data: bytes) -> dict[str, bool]:
    """Look for active or external PDF content without executing anything."""
    import re

    return {
        "javascript": bool(re.search(rb"/(JavaScript|JS)\b", data)),
        "embedded_files": bool(re.search(rb"/Type\s*/EmbeddedFile\b", data)),
        "forms": b"/AcroForm" in data,
        "signature": bool(re.search(rb"/Type\s*/Sig\b", data)) or (b"/ByteRange" in data and b"/Sig" in data),
        "external_links": bool(re.search(rb"/(URI|Launch|GoToR|SubmitForm)\b", data)),
    }


def analyze(data: bytes, view: str) -> dict[str, Any]:
    """Extract metadata, analyst signals and searchable text for one document type."""
    if view == "pdf":
        return _analyze_pdf(data)
    if view == "docx":
        return _analyze_docx(data)
    if view == "xlsx":
        return _analyze_xlsx(data)
    if view == "xls":
        return _analyze_xls(data)
    if view == "odf":
        return _analyze_odf(data)
    return {"metadata": {}, "signals": [], "text": None}


def _analyze_odf(data: bytes) -> dict[str, Any]:
    """OpenDocument properties (meta.xml) and macro, event and external-link signals."""
    import re
    import zipfile

    metadata: dict[str, str] = {}
    signals: list[str] = []
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        names = archive.namelist()
        if "meta.xml" in names:
            meta = archive.read("meta.xml").decode("utf-8", "replace")
            for tag, label in (
                ("dc:title", "Cím"), ("meta:initial-creator", "Szerző"), ("dc:creator", "Utolsó mentő"),
                ("meta:creation-date", "Létrehozva"), ("dc:date", "Módosítva"), ("meta:generator", "Alkalmazás"),
            ):
                found = re.search(rf"<{tag}>([^<]*)</{tag}>", meta)
                if found and found.group(1).strip():
                    value = found.group(1).strip()
                    metadata[label] = value.replace("T", " ")[:16] if tag.endswith("date") else value
        modules = [
            name for name in names
            if name.startswith("Basic/") and name.endswith(".xml") and not name.endswith(("script-lb.xml", "script-lc.xml"))
        ]
        scripts = [name for name in names if name.startswith("Scripts/") and not name.endswith("/")]
        if modules or scripts:
            signals.append(f"Makrót tartalmaz ({len(modules) + len(scripts)} modul, Basic/Scripts) – a nézet nem futtatja.")
        content = archive.read("content.xml") if "content.xml" in names else b""
        events = len(re.findall(rb"<script:event-listener\b", content))
        if events:
            signals.append(f"Eseményhez (pl. megnyitáshoz) kötött makró: {events} db.")
        external = len(re.findall(rb'xlink:href="(?:https?|ftp|file):', content))
        if external:
            signals.append(f"Külső hivatkozás: {external} db (az előnézet nem követi).")
    return {"metadata": metadata, "signals": signals}


def _pdf_date(value: str) -> str:
    """Turn a PDF date (D:YYYYMMDDHHmmSS) into YYYY-MM-DD HH:MM."""
    import re

    match = re.match(r"D?:?(\d{4})(\d{2})?(\d{2})?(\d{2})?(\d{2})?", value or "")
    if not match:
        return value
    year, month, day, hour, minute = (part or "00" for part in match.groups())
    return f"{year}-{month or '01'}-{day or '01'} {hour}:{minute}"


def _analyze_pdf(data: bytes) -> dict[str, Any]:
    import pypdfium2 as pdfium

    document = pdfium.PdfDocument(data)
    try:
        info = document.get_metadata_dict()
        version = document.get_version()
        pages = len(document)
        texts: list[str] = []
        total = 0
        for index in range(pages):
            if total > DOCUMENT_TEXT_LIMIT:
                texts.append("")
                continue
            text = document[index].get_textpage().get_text_range()
            texts.append(text)
            total += len(text.encode("utf-8"))
        try:
            attachments = document.count_attachments()
        except Exception:
            attachments = 0
    finally:
        document.close()
    signals = _pdf_signals(data)
    if attachments:
        signals["embedded_files"] = True
    has_text = any(text.strip() for text in texts)
    metadata = {
        "Cím": info.get("Title", ""),
        "Szerző": info.get("Author", ""),
        "Tárgy": info.get("Subject", ""),
        "Készítő program": info.get("Creator", ""),
        "PDF-előállító": info.get("Producer", ""),
        "Létrehozva": _pdf_date(info.get("CreationDate", "")),
        "Módosítva": _pdf_date(info.get("ModDate", "")),
        "PDF-verzió": f"{version / 10:.1f}" if version else "",
        "Oldalak": str(pages),
    }
    flagged = []
    if signals["javascript"]:
        flagged.append("Beágyazott JavaScript (az előnézet nem futtatja).")
    if signals["embedded_files"]:
        flagged.append("Beágyazott fájl a PDF-ben.")
    if signals["forms"]:
        flagged.append("Űrlapmezők vannak a dokumentumban.")
    if signals["signature"]:
        flagged.append("Digitális aláírás található (az aláírást a program nem ellenőrzi).")
    if signals["external_links"]:
        flagged.append("Külső hivatkozás vagy indítási művelet (az előnézet nem követi).")
    if not has_text:
        flagged.append("Nincs szövegréteg: a szöveg nem kinyerhető (OCR nélkül).")
    return {
        "metadata": {key: value for key, value in metadata.items() if value},
        "signals": flagged,
        "security": {
            "Beágyazott JavaScript": "van" if signals["javascript"] else "nincs",
            "Beágyazott fájl": "van" if signals["embedded_files"] else "nincs",
            "Űrlapmezők": "van" if signals["forms"] else "nincs",
            "Digitális aláírás": "van" if signals["signature"] else "nincs",
            "Külső hivatkozás": "van" if signals["external_links"] else "nincs",
        },
        "text_layer": has_text,
        "pages_text": texts,
    }


def _office_properties(data: bytes) -> dict[str, str]:
    """Read author, editor, dates and application from Office core.xml / app.xml."""
    import zipfile
    from xml.etree import ElementTree

    names = {
        "creator": "Szerző",
        "lastModifiedBy": "Utolsó mentő",
        "title": "Cím",
        "created": "Létrehozva",
        "modified": "Módosítva",
        "Application": "Alkalmazás",
        "Company": "Cég",
    }
    result: dict[str, str] = {}
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        for part in ("docProps/core.xml", "docProps/app.xml"):
            try:
                root = ElementTree.fromstring(archive.read(part))
            except (KeyError, ElementTree.ParseError):
                continue
            for element in root.iter():
                local = element.tag.rsplit("}", 1)[-1]
                if local in names and (element.text or "").strip():
                    value = element.text.strip()
                    if local in {"created", "modified"}:
                        value = value.replace("T", " ").rstrip("Z")[:16]
                    result[names[local]] = value
    return result


def _office_signals(data: bytes) -> list[str]:
    """Flag macros, external relationships, comments and embedded objects in OOXML."""
    import re
    import zipfile

    signals: list[str] = []
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        names = archive.namelist()
        if any(name.endswith("vbaProject.bin") for name in names):
            signals.append("Makrót tartalmaz (vbaProject.bin) – az előnézet nem futtatja.")
        external = 0
        for name in names:
            if name.endswith(".rels"):
                external += len(re.findall(rb'TargetMode="External"', archive.read(name)))
        if external:
            signals.append(f"Külső hivatkozás: {external} db (az előnézet nem követi).")
        embedded = [name for name in names if "/embeddings/" in name]
        if embedded:
            signals.append(f"Beágyazott objektum: {len(embedded)} db.")
        if "word/comments.xml" in names:
            comments = len(re.findall(rb"<w:comment\b", archive.read("word/comments.xml")))
            if comments:
                signals.append(f"Megjegyzés: {comments} db.")
    return signals


def render_docx(data: bytes, filename: str) -> dict[str, Any]:
    """Convert DOCX to sanitized HTML (formatted view) and plain text."""
    import mammoth

    converted = mammoth.convert_to_html(io.BytesIO(data))
    text = mammoth.extract_raw_text(io.BytesIO(data)).value
    text, truncated = _limited(text)
    return {
        "kind": "docx",
        "document": _document(sanitize_html(converted.value, keep_styles=False), filename),
        "text": text,
        "truncated": truncated,
        "messages": [str(message.message) for message in converted.messages][:20],
    }


def _analyze_docx(data: bytes) -> dict[str, Any]:
    return {"metadata": _office_properties(data), "signals": _office_signals(data)}


def _cell_text(value: Any) -> str:
    """Show a spreadsheet value the way an analyst expects to read it."""
    import datetime

    if value is None:
        return ""
    if isinstance(value, datetime.datetime):
        return value.strftime("%Y-%m-%d %H:%M") if (value.hour or value.minute) else value.strftime("%Y-%m-%d")
    if isinstance(value, datetime.date):
        return value.isoformat()
    if isinstance(value, float):
        return f"{value:.10g}"
    return str(value)


def _column_letter(index: int) -> str:
    letters = ""
    index += 1
    while index:
        index, remainder = divmod(index - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def render_xlsx(data: bytes, sheet: int = 0, page: int = 0) -> dict[str, Any]:
    """Return one sheet page: stored values, formulas, hidden rows/columns, comments."""
    from openpyxl import load_workbook

    formulas_book = load_workbook(io.BytesIO(data), data_only=False)
    values_book = load_workbook(io.BytesIO(data), data_only=True)
    sheets = [{"name": ws.title, "state": ws.sheet_state} for ws in formulas_book.worksheets]
    index = min(max(0, sheet), len(sheets) - 1)
    formula_sheet = formulas_book.worksheets[index]
    value_sheet = values_book.worksheets[index]
    max_row = formula_sheet.max_row or 0
    max_column = min(formula_sheet.max_column or 0, SHEET_MAX_COLUMNS)
    start = max(0, page) * SHEET_PAGE_ROWS
    rows: list[list[str]] = []
    formulas: dict[str, str] = {}
    comments: dict[str, str] = {}
    hidden_rows: list[int] = []
    for row_number in range(start + 1, min(max_row, start + SHEET_PAGE_ROWS) + 1):
        if formula_sheet.row_dimensions[row_number].hidden:
            hidden_rows.append(row_number)
        row: list[str] = []
        for column in range(1, max_column + 1):
            cell = formula_sheet.cell(row=row_number, column=column)
            coordinate = cell.coordinate
            raw = cell.value
            if isinstance(raw, str) and raw.startswith("="):
                formulas[coordinate] = raw
            elif raw is not None and cell.data_type == "f":
                formulas[coordinate] = str(raw)
            if cell.comment is not None:
                comments[coordinate] = f"{cell.comment.author or ''}: {cell.comment.text}".strip(": ")
            row.append(_cell_text(value_sheet.cell(row=row_number, column=column).value))
        rows.append(row)
    hidden_columns = [
        _column_letter(column - 1)
        for column in range(1, max_column + 1)
        if formula_sheet.column_dimensions[_column_letter(column - 1)].hidden
    ]
    return {
        "kind": "sheet",
        "sheets": sheets,
        "sheet": index,
        "page": max(0, page),
        "page_size": SHEET_PAGE_ROWS,
        "total_rows": max_row,
        "first_row": start + 1,
        "columns": [_column_letter(column) for column in range(max_column)],
        "rows": rows,
        "formulas": formulas,
        "comments": comments,
        "hidden_rows": hidden_rows,
        "hidden_columns": hidden_columns,
        "values_missing": bool(formulas) and all(
            value_sheet[coordinate].value is None for coordinate in formulas
        ),
    }


def _analyze_xlsx(data: bytes) -> dict[str, Any]:
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(data), data_only=False)
    hidden = [ws.title for ws in book.worksheets if ws.sheet_state == "hidden"]
    very_hidden = [ws.title for ws in book.worksheets if ws.sheet_state == "veryHidden"]
    formulas = comments = hidden_rows = hidden_columns = 0
    for ws in book.worksheets:
        hidden_rows += sum(1 for dimension in ws.row_dimensions.values() if dimension.hidden)
        hidden_columns += sum(1 for dimension in ws.column_dimensions.values() if dimension.hidden)
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
                    formulas += 1
                if cell.comment is not None:
                    comments += 1
    signals = _office_signals(data)
    if hidden:
        signals.append(f"Rejtett munkalap: {', '.join(hidden)}")
    if very_hidden:
        signals.append(f"Nagyon rejtett munkalap (az Excelben sem látható): {', '.join(very_hidden)}")
    if hidden_rows:
        signals.append(f"Rejtett sor: {hidden_rows} db")
    if hidden_columns:
        signals.append(f"Rejtett oszlop: {hidden_columns} db")
    external = len(getattr(book, "_external_links", []) or [])
    if external:
        signals.append(f"Külső munkafüzet-hivatkozás: {external} db")
    counts = {
        "Munkalapok": str(len(book.worksheets)),
        "Képletek": str(formulas),
        "Megjegyzések": str(comments),
    }
    return {"metadata": {**_office_properties(data), **counts}, "signals": signals}


def render_xls(data: bytes, sheet: int = 0, page: int = 0) -> dict[str, Any]:
    """Return one sheet page of a legacy XLS workbook (values only)."""
    import xlrd

    book = xlrd.open_workbook(file_contents=data, on_demand=True)
    states = {0: "visible", 1: "hidden", 2: "veryHidden"}
    sheets = [{"name": book.sheet_by_index(i).name, "state": states.get(book.sheet_by_index(i).visibility, "visible")}
              for i in range(book.nsheets)]
    index = min(max(0, sheet), book.nsheets - 1)
    ws = book.sheet_by_index(index)
    start = max(0, page) * SHEET_PAGE_ROWS
    max_column = min(ws.ncols, SHEET_MAX_COLUMNS)
    rows: list[list[str]] = []
    for row_number in range(start, min(ws.nrows, start + SHEET_PAGE_ROWS)):
        row: list[str] = []
        for column in range(max_column):
            cell = ws.cell(row_number, column)
            if cell.ctype == xlrd.XL_CELL_DATE:
                row.append(_cell_text(xlrd.xldate_as_datetime(cell.value, book.datemode)))
            elif cell.ctype == xlrd.XL_CELL_NUMBER:
                row.append(_cell_text(float(cell.value)))
            else:
                row.append(_cell_text(cell.value if cell.value != "" else None))
        rows.append(row)
    return {
        "kind": "sheet",
        "sheets": sheets,
        "sheet": index,
        "page": max(0, page),
        "page_size": SHEET_PAGE_ROWS,
        "total_rows": ws.nrows,
        "first_row": start + 1,
        "columns": [_column_letter(column) for column in range(max_column)],
        "rows": rows,
        "formulas": {},
        "comments": {},
        "hidden_rows": [],
        "hidden_columns": [],
        "values_missing": False,
        "legacy": True,
    }


def _analyze_xls(data: bytes) -> dict[str, Any]:
    import xlrd

    book = xlrd.open_workbook(file_contents=data, on_demand=True)
    hidden = [book.sheet_by_index(i).name for i in range(book.nsheets) if book.sheet_by_index(i).visibility == 1]
    very_hidden = [book.sheet_by_index(i).name for i in range(book.nsheets) if book.sheet_by_index(i).visibility == 2]
    signals = []
    if hidden:
        signals.append(f"Rejtett munkalap: {', '.join(hidden)}")
    if very_hidden:
        signals.append(f"Nagyon rejtett munkalap: {', '.join(very_hidden)}")
    signals.append("Régi XLS-formátum: képletek és megjegyzések csak a „Pontos nézet”-ben (LibreOffice) látszanak.")
    metadata = {"Munkalapok": str(book.nsheets)}
    if book.user_name:
        metadata["Utolsó mentő"] = book.user_name
    return {"metadata": metadata, "signals": signals}

